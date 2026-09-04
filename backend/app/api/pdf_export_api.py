"""
PDF Export API — Generate printable reports for portfolios and compliance.
"""
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import HTMLResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import User

router = APIRouter(prefix="/pdf-export", tags=["pdf-export"])


def _get_user(request: Request, db: Session):
    token = request.cookies.get("access_token", "")
    if not token:
        return None
    try:
        from app.security import decode_token
        payload = decode_token(token)
        if payload:
            return db.query(User).filter(User.id == payload.get("sub")).first()
    except Exception:
        pass
    return None


class PortfolioReportRequest(BaseModel):
    portfolio_id: Optional[str] = None


@router.post("/portfolio", response_class=HTMLResponse)
def export_portfolio_report(
    request: Request,
    data: PortfolioReportRequest,
    db: Session = Depends(get_db),
):
    """Generate a printable portfolio report."""
    from app.services.pdf_export import generate_portfolio_report

    user = _get_user(request, db)
    if not user:
        raise HTTPException(status_code=401, detail="Not authenticated")

    from app.models import Portfolio
    from app.models import DebtInstrument

    portfolios = db.query(Portfolio).filter(Portfolio.org_id == user.org_id).all()
    portfolio_ids = [p.id for p in portfolios]
    instruments = db.query(DebtInstrument).filter(
        DebtInstrument.portfolio_id.in_(portfolio_ids)
    ).all() if portfolio_ids else []

    inst_list = []
    for i in instruments:
        inst_list.append({
            "name": i.name,
            "instrument_type": i.instrument_type,
            "principal_outstanding": float(i.principal_outstanding),
            "currency": i.currency,
            "coupon_rate": float(i.coupon_rate) if i.coupon_rate else 0,
            "maturity_date": str(i.maturity_date) if i.maturity_date else None,
        })

    portfolio_name = "Sovereign Debt Portfolio"
    if data.portfolio_id:
        for p in portfolios:
            if str(p.id) == data.portfolio_id:
                portfolio_name = p.name
                break

    html = generate_portfolio_report(inst_list, portfolio_name)
    return HTMLResponse(content=html)


@router.post("/compliance", response_class=HTMLResponse)
def export_compliance_report(
    request: Request,
    data: PortfolioReportRequest,
    db: Session = Depends(get_db),
):
    """Generate a printable compliance report."""
    from app.services.pdf_export import generate_compliance_report
    from app.services.compliance_checker import check_compliance

    user = _get_user(request, db)
    if not user:
        raise HTTPException(status_code=401, detail="Not authenticated")

    from app.models import Portfolio
    from app.models import DebtInstrument

    portfolios = db.query(Portfolio).filter(Portfolio.org_id == user.org_id).all()
    portfolio_ids = [p.id for p in portfolios]
    instruments = db.query(DebtInstrument).filter(
        DebtInstrument.portfolio_id.in_(portfolio_ids)
    ).all() if portfolio_ids else []

    inst_list = []
    for i in instruments:
        inst_list.append({
            "name": i.name,
            "instrument_type": i.instrument_type,
            "principal_outstanding": float(i.principal_outstanding),
            "currency": i.currency,
            "coupon_rate": float(i.coupon_rate) if i.coupon_rate else 0,
            "maturity_date": str(i.maturity_date) if i.maturity_date else None,
            "is_callable": getattr(i, "is_callable", False),
        })

    compliance_result = check_compliance(inst_list)
    html = generate_compliance_report(compliance_result)
    return HTMLResponse(content=html)


@router.post("/portfolio/excel")
def export_portfolio_excel(
    request: Request,
    data: PortfolioReportRequest,
    db: Session = Depends(get_db),
):
    """Generate Excel export of portfolio data."""
    from fastapi.responses import StreamingResponse
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    user = _get_user(request, db)
    if not user:
        raise HTTPException(status_code=401, detail="Not authenticated")

    from app.models import Portfolio
    from app.models import DebtInstrument

    portfolios = db.query(Portfolio).filter(Portfolio.org_id == user.org_id).all()
    portfolio_ids = [p.id for p in portfolios]
    instruments = db.query(DebtInstrument).filter(
        DebtInstrument.portfolio_id.in_(portfolio_ids)
    ).all() if portfolio_ids else []

    wb = Workbook()

    # ── Instruments Sheet ──
    ws = wb.active
    ws.title = "Instruments"
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        bottom=Side(style="thin", color="CCCCCC")
    )

    headers = ["#", "Name", "Type", "Currency", "Principal", "Coupon", "Maturity", "Issue Date", "Callable"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    total = 0
    for i, inst in enumerate(instruments, 2):
        principal = float(inst.principal_outstanding)
        total += principal
        ws.cell(row=i, column=1, value=i-1)
        ws.cell(row=i, column=2, value=inst.name)
        ws.cell(row=i, column=3, value=inst.instrument_type)
        ws.cell(row=i, column=4, value=inst.currency)
        ws.cell(row=i, column=5, value=principal).number_format = '#,##0'
        ws.cell(row=i, column=6, value=float(inst.coupon_rate) if inst.coupon_rate else 0).number_format = '0.00%'
        ws.cell(row=i, column=7, value=str(inst.maturity_date)[:10] if inst.maturity_date else "")
        ws.cell(row=i, column=8, value=str(inst.issue_date)[:10] if inst.issue_date else "")
        ws.cell(row=i, column=9, value="Yes" if getattr(inst, "is_callable", False) else "No")
        for col in range(1, 10):
            ws.cell(row=i, column=col).border = thin_border

    # Summary row
    summary_row = len(instruments) + 3
    ws.cell(row=summary_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=summary_row, column=5, value=total).number_format = '#,##0'
    ws.cell(row=summary_row, column=5).font = Font(bold=True)
    ws.cell(row=summary_row, column=2, value=f"{len(instruments)} instruments")

    # Auto-width columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 30)

    # ── Currency Breakdown Sheet ──
    ws2 = wb.create_sheet("Currency Breakdown")
    currencies = {}
    for inst in instruments:
        cur = inst.currency
        currencies[cur] = currencies.get(cur, 0) + float(inst.principal_outstanding)

    ws2.cell(row=1, column=1, value="Currency").font = header_font
    ws2.cell(row=1, column=1).fill = header_fill
    ws2.cell(row=1, column=2, value="Amount").font = header_font
    ws2.cell(row=1, column=2).fill = header_fill
    ws2.cell(row=1, column=3, value="Weight %").font = header_font
    ws2.cell(row=1, column=3).fill = header_fill

    for i, (cur, val) in enumerate(sorted(currencies.items(), key=lambda x: x[1], reverse=True), 2):
        ws2.cell(row=i, column=1, value=cur)
        ws2.cell(row=i, column=2, value=val).number_format = '#,##0'
        ws2.cell(row=i, column=3, value=round(val/total*100, 1) if total else 0).number_format = '0.0"%"'

    # ── Metadata Sheet ──
    ws3 = wb.create_sheet("Report Info")
    ws3.cell(row=1, column=1, value="Quantive Portfolio Export")
    ws3.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws3.cell(row=3, column=1, value="Generated:")
    ws3.cell(row=3, column=2, value=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
    ws3.cell(row=4, column=1, value="User:")
    ws3.cell(row=4, column=2, value=user.email)
    ws3.cell(row=5, column=1, value="Instruments:")
    ws3.cell(row=5, column=2, value=len(instruments))
    ws3.cell(row=6, column=1, value="Total Principal:")
    ws3.cell(row=6, column=2, value=total).number_format = '#,##0'
    ws3.cell(row=8, column=1, value="DISCLAIMER: FOR DECISION-SUPPORT PURPOSES ONLY.")
    ws3.cell(row=8, column=1).font = Font(italic=True, color="FF0000")

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"quantive_portfolio_{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
